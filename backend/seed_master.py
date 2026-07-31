"""
seed_master.py — Upsert tabel pendaftar dari Excel/CSV.
Pakai:  python seed_master.py for_seed.xlsx --dry-run
        python seed_master.py for_seed.xlsx
"""
import asyncio, os, re, sys
from pathlib import Path

BERKAS = next((a for a in sys.argv[1:] if not a.startswith("--")), "for_seed.xlsx")
DRY_RUN = "--dry-run" in sys.argv
DATABASE_URL = os.getenv(
    "DATABASE_URL", "postgresql+asyncpg://admin:secretpassword@db:5432/compliance_db"
)

def norm(s: object) -> str:
    return re.sub(r"[^a-z]", "", str(s or "").lower())

def baca_baris(path: str) -> list[list]:
    p = Path(path)
    if not p.exists():
        sys.exit(f"[X] Berkas tidak ditemukan: {p.resolve()}")
    if p.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        ws = load_workbook(p, data_only=True).active
        return [list(r) for r in ws.iter_rows(values_only=True)]
    import csv
    with open(p, encoding="utf-8-sig", newline="") as f:
        contoh = f.read(8192); f.seek(0)
        pemisah = ";" if contoh.count(";") > contoh.count(",") else ","
        return [r for r in csv.reader(f, delimiter=pemisah)]

def cari_header(baris: list[list]) -> tuple[int, dict]:
    """Cari baris header + indeks kolom id / nama / jurusan (fleksibel)."""
    for i, r in enumerate(baris[:10]):
        kol = {}
        for j, sel in enumerate(r):
            n = norm(sel)
            if not n:
                continue
            if "nama" in n and "nama" not in kol:
                kol["nama"] = j
            elif ("jurusan" in n or "prodi" in n or "program" in n) and "jurusan" not in kol:
                kol["jurusan"] = j
            elif ("id" in n or "nomor" in n or "no" == n) and "id" not in kol:
                kol["id"] = j
        if {"id", "nama", "jurusan"} <= kol.keys():
            return i, kol
    sys.exit("[X] Header tidak dikenali. Butuh kolom mengandung: id, nama, jurusan.")

def kumpulkan(baris, mulai, kol):
    sah, ditolak, terlihat = [], [], {}
    for n, r in enumerate(baris[mulai + 1:], start=mulai + 2):
        ambil = lambda k: str(r[kol[k]]).strip() if kol[k] < len(r) and r[kol[k]] is not None else ""
        idp, nama, jur = ambil("id"), ambil("nama"), ambil("jurusan")
        idp = idp.split(".")[0]  # Excel kadang membaca angka jadi 426161036.0
        if not any((idp, nama, jur)):
            continue
        if not re.fullmatch(r"\d{4,}", idp):
            ditolak.append((n, idp, "id_pendaftaran bukan angka")); continue
        if not nama or nama.lower() in ("none", "nan", "#n/a"):
            ditolak.append((n, idp, "nama kosong")); continue
        if not jur:
            ditolak.append((n, idp, "jurusan kosong")); continue
        if idp in terlihat:
            ditolak.append((n, idp, f"duplikat (sudah di baris {terlihat[idp]})")); continue
        terlihat[idp] = n
        sah.append({"idp": idp, "nama": " ".join(nama.split()), "jur": " ".join(jur.split())})
    return sah, ditolak

SQL_UPSERT = """
INSERT INTO pendaftar (id_pendaftaran, nama, jurusan_dituju, created_at)
VALUES (:idp, :nama, :jur, NOW())
ON CONFLICT (id_pendaftaran) DO UPDATE
   SET nama = EXCLUDED.nama,
       jurusan_dituju = EXCLUDED.jurusan_dituju
"""

async def main():
    baris = baca_baris(BERKAS)
    i, kol = cari_header(baris)
    sah, ditolak = kumpulkan(baris, i, kol)

    print(f"Berkas   : {BERKAS}")
    print(f"Header   : baris {i+1} -> {kol}")
    print(f"Sah      : {len(sah)} baris")
    print(f"Ditolak  : {len(ditolak)} baris")
    for n, idp, sebab in ditolak:
        print(f"   - baris {n} ({idp or 'kosong'}): {sebab}")
    if sah[:3]:
        print("Contoh   :", *[f"\n   {r['idp']} | {r['nama']} | {r['jur']}" for r in sah[:3]])

    if DRY_RUN:
        print("\n[DRY-RUN] Basis data tidak disentuh.")
        return
    if not sah:
        sys.exit("[X] Tidak ada baris sah. Batal.")

    from sqlalchemy import text
    from sqlalchemy.ext.asyncio import create_async_engine
    engine = create_async_engine(DATABASE_URL)
    async with engine.begin() as conn:
        sebelum = (await conn.execute(text("SELECT count(*) FROM pendaftar"))).scalar()
        for r in sah:
            await conn.execute(text(SQL_UPSERT), r)
        sesudah = (await conn.execute(text("SELECT count(*) FROM pendaftar"))).scalar()
        kosong = (await conn.execute(text(
            "SELECT count(*) FROM pendaftar WHERE nama IS NULL OR nama = ''"))).scalar()
    await engine.dispose()

    print(f"\nSelesai  : {sebelum} -> {sesudah} baris (baru {sesudah - sebelum}, diperbarui {len(sah) - (sesudah - sebelum)})")
    print(f"Nama kosong di tabel: {kosong}  {'[OK]' if kosong == 0 else '[PERIKSA!]'}")

if __name__ == "__main__":
    asyncio.run(main())