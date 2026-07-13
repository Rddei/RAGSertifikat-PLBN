#!/usr/bin/env python3
"""
Scraper SIMT PUSPRESNAS - Daftar Ajang Talenta Terkurasi (SEMUA data).

Tujuan:
    Mengambil SELURUH daftar ajang terkurasi dari halaman
    https://simt.kemendikdasmen.go.id/kurasi/advanced-search
    (± 6.035 baris / 604 halaman) menjadi snapshot LOKAL, sehingga engine
    verifikasi tidak perlu live-scraping tiap sertifikat (lebih cepat & andal).

Kenapa Playwright (headless browser), bukan httpx?
    Halaman advanced-search me-render baris pertama secara server-side, TAPI
    navigasi halaman ("1 2 3 ... 604 Next") dilakukan lewat JavaScript/AJAX
    (tombol nomor tidak punya href, dan ?page=N tidak mengembalikan data).
    Jadi kita butuh browser sungguhan yang meng-klik "Next".

Output (default folder ./data_simt/):
    - simt_ajang.xlsx    -> spreadsheet Excel siap ditinjau (DELIVERABLE UTAMA)
    - simt_ajang.jsonl   -> 1 baris JSON per ajang (append, dipakai untuk --resume)
    - .progress          -> checkpoint halaman terakhir yang selesai

Catatan: skrip ini SENGAJA hanya menghasilkan file (belum menyentuh database/
engine). Tinjau dulu .xlsx-nya; integrasi ke sistem menyusul terpisah.

Cara pakai:
    pip install playwright
    playwright install chromium
    python scrape_simt.py                 # scrape semua, headless
    python scrape_simt.py --headed        # tampilkan browser (debug)
    python scrape_simt.py --max-pages 5   # uji cepat 5 halaman dulu
    python scrape_simt.py --delay 0.8     # jeda antar halaman (detik)
    python scrape_simt.py --resume        # lanjutkan dari checkpoint

Catatan:
    - Skrip TIDAK melempar error fatal di tengah jalan: tiap halaman disimpan
      segera (append JSONL), jadi kalau putus bisa --resume.
    - Deduplikasi berdasarkan SELURUH kolom (termasuk nomor urut "no"), sehingga
      baris yang identik kontennya namun sah tetap tersimpan; hanya halaman yang
      benar-benar ter-ekstrak ulang yang ter-dedup.
    - Struktur DOM situs bisa berubah sewaktu-waktu; selector "Next" dibuat
      defensif (beberapa kandidat). Kalau situs berubah total, jalankan
      --headed untuk melihat & sesuaikan _click_next().
"""
from __future__ import annotations

import argparse
import asyncio
import json
import logging
import re
import sys
from pathlib import Path
from typing import Any

try:
    from playwright.async_api import async_playwright, TimeoutError as PWTimeout
except ImportError:  # pragma: no cover
    sys.exit(
        "Playwright belum terpasang. Jalankan:\n"
        "    pip install playwright && playwright install chromium"
    )

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s  %(levelname)-7s %(message)s",
    datefmt="%H:%M:%S",
)
log = logging.getLogger("scrape_simt")

URL = "https://simt.kemendikdasmen.go.id/kurasi/advanced-search"

# Peta header tabel -> key JSON. Dicocokkan longgar (lowercase, tanpa spasi ganda).
HEADER_MAP = {
    "no": "no",
    "cabang": "cabang",
    "nama ajang": "nama_ajang",
    "hasil kurasi": "hasil_kurasi",
    "singkatan": "singkatan",
    "penyelenggara": "penyelenggara",
    "negara": "negara",
    "tipe penyelenggaraan": "tipe",
    "level": "level",
    "kategori": "kategori",
    "tanggal mulai": "tanggal_mulai",
    "tanggal selesai": "tanggal_selesai",
    "link": "link",
}

_RANGE_RE = re.compile(
    r"Menampilkan\s+([\d.,]+)\s*-\s*([\d.,]+)\s+dari\s+([\d.,]+)\s+data",
    re.IGNORECASE,
)

# JS yang mengekstrak isi tabel + header + teks paginasi dari halaman aktif.
# Dikembalikan sebagai objek {headers, rows, rangeText} agar parsing terpusat.
EXTRACT_JS = r"""
() => {
  const table = document.querySelector('table');
  if (!table) return { headers: [], rows: [], rangeText: '' };

  const norm = (s) => (s || '').replace(/\s+/g, ' ').trim();

  const headers = [...table.querySelectorAll('thead th')].map(th => norm(th.innerText));

  const bodyRows = [...table.querySelectorAll('tbody tr')];
  const rows = bodyRows.map(tr => {
    return [...tr.children].map(td => {
      const a = td.querySelector('a[href]');
      // Hitung bintang "terisi" secara best-effort dari ikon di dalam sel.
      let starsFilled = 0, starsTotal = 0;
      td.querySelectorAll('svg, i, span, img').forEach(el => {
        let cls = '';
        if (el.className && typeof el.className === 'object' && 'baseVal' in el.className) {
          cls = el.className.baseVal;              // SVG
        } else {
          cls = String(el.className || '');
        }
        cls = cls.toLowerCase();
        const looksLikeStar = cls.includes('star') || cls.includes('fa-star') || cls.includes('rating');
        if (!looksLikeStar) return;
        starsTotal++;
        const filled = (cls.includes('fill') || cls.includes('active') ||
                        cls.includes('checked') || cls.includes('warning') ||
                        cls.includes('fas') || cls.includes('text-yellow') ||
                        cls.includes('text-warning')) &&
                       !cls.includes('-o') && !cls.includes('far') && !cls.includes('empty');
        if (filled) starsFilled++;
      });
      return {
        text: norm(td.innerText),
        href: a ? a.href : null,
        starsFilled: starsFilled,
        starsTotal: starsTotal,
      };
    });
  });

  const bodyText = norm(document.body.innerText);
  return { headers, rows, rangeText: bodyText };
}
"""


def _parse_int(s: str) -> int | None:
    digits = re.sub(r"[^\d]", "", s or "")
    return int(digits) if digits else None


def _parse_range(text: str) -> tuple[int | None, int | None, int | None]:
    """Ambil (start, end, total) dari 'Menampilkan a - b dari N data'."""
    m = _RANGE_RE.search(text or "")
    if not m:
        return None, None, None
    return _parse_int(m.group(1)), _parse_int(m.group(2)), _parse_int(m.group(3))


def _row_to_dict(cells: list[dict], header_keys: list[str]) -> dict[str, Any]:
    """Petakan sel-sel satu baris ke dict berdasarkan urutan header."""
    row: dict[str, Any] = {}
    for idx, cell in enumerate(cells):
        key = header_keys[idx] if idx < len(header_keys) else f"col{idx}"
        if key == "link":
            row["link"] = cell.get("href") or (cell.get("text") or None)
        elif key == "hasil_kurasi":
            # Utamakan angka bila ada (mis. "5.0"); jika tidak, pakai hitungan bintang.
            txt = cell.get("text") or ""
            num = re.search(r"\d+(?:[.,]\d+)?", txt)
            if num:
                row["hasil_kurasi"] = float(num.group(0).replace(",", "."))
            elif cell.get("starsFilled"):
                row["hasil_kurasi"] = float(cell["starsFilled"])
            else:
                row["hasil_kurasi"] = None
        else:
            row[key] = cell.get("text") or None
    return row


def _dedup_key(row: dict[str, Any]) -> str:
    def n(v):
        return re.sub(r"\s+", " ", str(v if v is not None else "").strip().lower())
    # Kunci dedup memakai SEMUA kolom, termasuk nomor urut "no". Situs SIMT
    # ternyata memuat baris-baris yang identik di semua kolom konten (beda hanya
    # pada "no"), jadi "no" WAJIB ikut agar baris sah tidak ikut terbuang.
    # Hanya penanda internal "_page" yang dikecualikan (bukan data situs, dan
    # baru ditambahkan setelah kunci dihitung -- penting untuk konsistensi resume).
    # Baris yang benar-benar sama (mis. halaman ter-ekstrak ulang) tetap ter-dedup
    # karena seluruh nilainya -- termasuk "no" -- juga sama persis.
    keys = sorted(k for k in row.keys() if k != "_page")
    return "|".join(f"{k}={n(row.get(k))}" for k in keys)


async def _click_next(page) -> bool:
    """Klik tombol/next paginasi. Kembalikan True jika kemungkinan berhasil.

    Beberapa kandidat selector dicoba (situs bisa berubah). Ambil yang PALING
    BAWAH (pagination biasanya di kaki tabel) untuk menghindari salah-klik.
    """
    candidates = [
        "xpath=//a[normalize-space(.)='Next']",
        "xpath=//button[normalize-space(.)='Next']",
        "xpath=//li[not(contains(@class,'disabled'))]/a[normalize-space(.)='Next']",
        "xpath=//a[normalize-space(.)='›' or normalize-space(.)='>']",
        "text=Next",
    ]
    for sel in candidates:
        try:
            loc = page.locator(sel).last
            if await loc.count() == 0:
                continue
            if not await loc.is_visible():
                continue
            if not await loc.is_enabled():
                continue
            await loc.scroll_into_view_if_needed(timeout=3000)
            await loc.click(timeout=5000)
            return True
        except Exception:
            continue
    return False


async def _extract(page) -> dict[str, Any]:
    return await page.evaluate(EXTRACT_JS)


async def scrape(args) -> None:
    out_dir = Path(args.out).resolve()
    out_dir.mkdir(parents=True, exist_ok=True)
    jsonl_path = out_dir / "simt_ajang.jsonl"
    progress_path = out_dir / ".progress"
    xlsx_path = out_dir / "simt_ajang.xlsx"

    seen: set[str] = set()
    resume_from = 0
    if args.resume and jsonl_path.exists():
        with jsonl_path.open(encoding="utf-8") as f:
            for line in f:
                try:
                    seen.add(_dedup_key(json.loads(line)))
                except Exception:
                    pass
        if progress_path.exists():
            resume_from = _parse_int(progress_path.read_text()) or 0
        log.info("Resume: %d ajang sudah ada, checkpoint halaman %d", len(seen), resume_from)
    else:
        # Mulai bersih.
        if jsonl_path.exists():
            jsonl_path.unlink()

    total_expected: int | None = None
    jsonl_f = jsonl_path.open("a", encoding="utf-8")

    async with async_playwright() as pw:
        launch_kwargs: dict[str, Any] = {"headless": not args.headed}
        if args.chromium_path:
            # Pakai Chromium sistem (mis. omarchy-chromium di /usr/bin/chromium)
            # alih-alih Chromium bawaan Playwright.
            launch_kwargs["executable_path"] = args.chromium_path
        browser = await pw.chromium.launch(**launch_kwargs)
        ctx = await browser.new_context(
            user_agent=(
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                "(KHTML, like Gecko) Chrome/124.0 Safari/537.36"
            ),
            locale="id-ID",
        )
        page = await ctx.new_page()
        log.info("Membuka %s", URL)
        await page.goto(URL, wait_until="networkidle", timeout=60000)
        await page.wait_for_selector("table tbody tr", timeout=30000)

        raw = await _extract(page)
        header_keys = [HEADER_MAP.get(h.lower(), h.lower().replace(" ", "_"))
                       for h in raw["headers"]]
        log.info("Header terdeteksi: %s", header_keys)

        page_idx = 1
        # Fast-forward saat resume (klik Next tanpa memproses hingga checkpoint).
        while page_idx <= resume_from:
            if not await _click_next(page):
                break
            await page.wait_for_timeout(400)
            page_idx += 1
        if resume_from:
            log.info("Fast-forward selesai di halaman %d", page_idx)

        last_range_end = -1
        stagnant = 0

        while True:
            if args.max_pages and page_idx > (resume_from + args.max_pages):
                log.info("Berhenti: mencapai --max-pages", )
                break

            raw = await _extract(page)
            start, end, total = _parse_range(raw["rangeText"])
            if total:
                total_expected = total

            new_in_page = 0
            for cells in raw["rows"]:
                row = _row_to_dict(cells, header_keys)
                # Lewati baris kosong / baris "tidak ada data".
                if not (row.get("nama_ajang") or row.get("penyelenggara")):
                    continue
                key = _dedup_key(row)
                if key in seen:
                    continue
                seen.add(key)
                row["_page"] = page_idx
                jsonl_f.write(json.dumps(row, ensure_ascii=False) + "\n")
                new_in_page += 1
            jsonl_f.flush()
            progress_path.write_text(str(page_idx))

            log.info(
                "Halaman %s | rentang %s-%s dari %s | +%d baru | total tersimpan %d",
                page_idx, start, end, total_expected, new_in_page, len(seen),
            )

            # Kondisi berhenti: sudah mencapai baris terakhir.
            if total_expected and end and end >= total_expected:
                log.info("Selesai: mencapai baris terakhir (%s).", total_expected)
                break

            # Deteksi kemacetan (Next tidak mengubah rentang).
            if end is not None and end == last_range_end:
                stagnant += 1
            else:
                stagnant = 0
            last_range_end = end if end is not None else last_range_end
            if stagnant >= 3:
                log.warning("Berhenti: rentang tidak berubah 3x (mungkin akhir/DOM berubah).")
                break

            if not await _click_next(page):
                log.warning("Tombol 'Next' tidak ditemukan/klik gagal. Berhenti.")
                break

            # Tunggu rentang berubah (konten halaman ter-update via AJAX).
            try:
                await page.wait_for_function(
                    "(prevEnd) => {"
                    "  const m = document.body.innerText.match(/Menampilkan\\s+[\\d.,]+\\s*-\\s*([\\d.,]+)\\s+dari/i);"
                    "  if (!m) return false;"
                    "  const cur = parseInt(m[1].replace(/[^0-9]/g,''), 10);"
                    "  return cur !== prevEnd;"
                    "}",
                    arg=end,
                    timeout=15000,
                )
            except PWTimeout:
                log.warning("Timeout menunggu halaman berganti (halaman %d).", page_idx + 1)
            await page.wait_for_timeout(int(args.delay * 1000))
            page_idx += 1

        await browser.close()

    jsonl_f.close()
    log.info("Total ajang tersimpan: %d (perkiraan situs: %s)", len(seen), total_expected)

    _build_xlsx(jsonl_path, xlsx_path)
    log.info("XLSX   -> %s", xlsx_path)
    log.info("JSONL  -> %s (untuk --resume)", jsonl_path)


COLUMNS = [
    "no", "cabang", "nama_ajang", "hasil_kurasi", "singkatan", "penyelenggara",
    "negara", "tipe", "level", "kategori", "tanggal_mulai", "tanggal_selesai",
    "link", "_page",
]


def _iter_rows(jsonl_path: Path):
    with jsonl_path.open(encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if line:
                try:
                    yield json.loads(line)
                except Exception:
                    continue


def _build_xlsx(jsonl_path: Path, xlsx_path: Path) -> None:
    """Bangun spreadsheet Excel yang rapi dari JSONL hasil scrape."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    # (key JSON, label kolom, lebar)
    cols = [
        ("no", "No", 6),
        ("cabang", "Cabang", 30),
        ("nama_ajang", "Nama Ajang", 42),
        ("singkatan", "Singkatan", 14),
        ("penyelenggara", "Penyelenggara", 34),
        ("negara", "Negara", 12),
        ("tipe", "Tipe", 16),
        ("level", "Level", 15),
        ("kategori", "Kategori", 16),
        ("hasil_kurasi", "Hasil Kurasi", 12),
        ("tanggal_mulai", "Tanggal Mulai", 15),
        ("tanggal_selesai", "Tanggal Selesai", 15),
        ("link", "Link", 40),
        ("_page", "Halaman", 9),
    ]

    wb = Workbook()
    ws = wb.active
    ws.title = "Ajang Terkurasi SIMT"

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF")
    for c_idx, (_key, label, width) in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=c_idx, value=label)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c_idx)].width = width

    link_font = Font(color="0563C1", underline="single")
    r_idx = 2
    for row in _iter_rows(jsonl_path):
        for c_idx, (key, _label, _w) in enumerate(cols, start=1):
            val = row.get(key)
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.alignment = Alignment(vertical="top", wrap_text=(key in {"nama_ajang", "cabang", "penyelenggara"}))
            # Jadikan hyperlink hanya bila satu URL bersih (kadang ada multi-URL dipisah koma).
            if key == "link" and isinstance(val, str) and val.startswith("http") and " " not in val:
                cell.hyperlink = val
                cell.font = link_font
        r_idx += 1

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(r_idx - 1, 1)}"
    wb.save(xlsx_path)


def main() -> None:
    p = argparse.ArgumentParser(description="Scrape semua ajang terkurasi SIMT.")
    p.add_argument("--out", default="data_simt", help="Folder output (default: data_simt)")
    p.add_argument("--headed", action="store_true", help="Tampilkan browser (debug)")
    p.add_argument("--delay", type=float, default=0.6, help="Jeda antar halaman (detik)")
    p.add_argument("--max-pages", type=int, default=0, help="Batasi jumlah halaman (0=semua)")
    p.add_argument("--resume", action="store_true", help="Lanjut dari checkpoint")
    p.add_argument(
        "--chromium-path",
        default=None,
        help=("Path ke Chromium sistem, mis. /usr/bin/chromium (Omarchy/Arch). "
              "Default: pakai Chromium bawaan Playwright."),
    )
    args = p.parse_args()
    asyncio.run(scrape(args))


if __name__ == "__main__":
    main()
