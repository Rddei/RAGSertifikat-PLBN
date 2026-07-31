"""
Layanan data master PENDAFTAR + parser nama file sertifikat.

Alur input baru (permintaan pembimbing): verifikator cukup mengunggah file;
nama file memuat ID pendaftaran dengan format:  <id_pendaftar>-<indeks>.<ext>
contoh: 221524015-1.pdf  ->  id '221524015', sertifikat ke-1 (maks. 3).

Nama pendaftar dan jurusan tujuan TIDAK lagi diketik manual, melainkan
di-lookup dari tabel master `pendaftar` yang diimpor dari file CSV/XLSX.
File yang id-nya tidak terdaftar DITOLAK di awal ("pendaftar tidak tersedia").
"""

import csv
import io
import logging
import re

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from models.database import Pendaftar

logger = logging.getLogger("compliance.pendaftar")

MAX_INDEKS_SERTIFIKAT = 3

# <digit 4+> - <indeks 1..2 digit> [pemisah opsional sisa nama] . ext
_FILENAME_PAT = re.compile(
    r"^\s*(\d{4,})\s*-\s*(\d{1,2})(?:[\s._-].*)?\.(pdf|jpe?g|png|webp)\s*$",
    re.IGNORECASE,
)


def parse_certificate_filename(filename: str) -> tuple[str, int]:
    """Ambil (id_pendaftaran, indeks) dari nama file; ValueError bila tak sesuai.

    Contoh valid : '221524015-1.pdf', '221524015-2 revisi.jpg'
    Tidak valid  : 'sertifikat.pdf' (tanpa id), '221524015-4.pdf' (indeks > 3)
    """
    m = _FILENAME_PAT.match(filename or "")
    if not m:
        raise ValueError(
            f"Nama file '{filename}' tidak sesuai format '<id_pendaftar>-<indeks>.<ext>' "
            f"(contoh: 221524015-1.pdf)"
        )
    id_pendaftaran, indeks = m.group(1), int(m.group(2))
    if not (1 <= indeks <= MAX_INDEKS_SERTIFIKAT):
        raise ValueError(
            f"Indeks sertifikat pada '{filename}' adalah {indeks}; "
            f"maksimum {MAX_INDEKS_SERTIFIKAT} sertifikat per pendaftar"
        )
    return id_pendaftaran, indeks


async def get_pendaftar(db: AsyncSession, id_pendaftaran: str) -> Pendaftar | None:
    res = await db.execute(
        select(Pendaftar).where(Pendaftar.id_pendaftaran == id_pendaftaran)
    )
    return res.scalar_one_or_none()


# ---------------------------------------------------------------------------
# Impor master pendaftar dari CSV / XLSX.
# Header dikenali fleksibel: kolom mengandung 'id' -> id_pendaftaran,
# 'nama' -> nama, 'jurusan'/'prodi' -> jurusan_dituju.
# ---------------------------------------------------------------------------
def _map_headers(headers: list[str]) -> dict[str, int]:
    idx: dict[str, int] = {}
    for i, h in enumerate(headers):
        h_low = (h or "").strip().lower()
        if "id" in h_low and "id_pendaftaran" not in idx:
            idx["id_pendaftaran"] = i
        elif "nama" in h_low and "nama" not in idx:
            idx["nama"] = i
        elif ("jurusan" in h_low or "prodi" in h_low) and "jurusan" not in idx:
            idx["jurusan"] = i
    missing = {"id_pendaftaran", "nama", "jurusan"} - set(idx)
    if missing:
        raise ValueError(
            f"Header file master tidak lengkap; kolom tidak ditemukan: "
            f"{', '.join(sorted(missing))}. Header terbaca: {headers}"
        )
    return idx


def _rows_from_csv(data: bytes) -> list[list[str]]:
    text = data.decode("utf-8-sig", errors="replace")
    # Deteksi delimiter sederhana (koma/semicolon -- Excel lokal sering ';')
    delim = ";" if text.splitlines()[0].count(";") > text.splitlines()[0].count(",") else ","
    return [row for row in csv.reader(io.StringIO(text), delimiter=delim)]


def _rows_from_xlsx(data: bytes) -> list[list[str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError(
            "Membaca .xlsx membutuhkan paket 'openpyxl' "
            "(pip install openpyxl), atau unggah dalam format CSV."
        ) from exc
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.active
    return [["" if c is None else str(c) for c in row]
            for row in ws.iter_rows(values_only=True)]


async def import_master(db: AsyncSession, data: bytes, filename: str) -> dict:
    """Impor/​perbarui master pendaftar. Upsert berdasarkan id_pendaftaran."""
    name_low = (filename or "").lower()
    if name_low.endswith(".csv"):
        rows = _rows_from_csv(data)
    elif name_low.endswith((".xlsx", ".xlsm")):
        rows = _rows_from_xlsx(data)
    else:
        raise ValueError("Format file master harus .csv atau .xlsx")

    rows = [r for r in rows if any(str(c).strip() for c in r)]
    if len(rows) < 2:
        raise ValueError("File master kosong atau hanya berisi header")

    idx = _map_headers([str(c) for c in rows[0]])
    dibuat, diperbarui, dilewati = 0, 0, []
    for n, row in enumerate(rows[1:], start=2):
        try:
            idp = str(row[idx["id_pendaftaran"]]).strip()
            # Excel kerap membaca id numerik sebagai float ('221524015.0')
            idp = re.sub(r"\.0$", "", idp)
            nama = str(row[idx["nama"]]).strip()
            jurusan = str(row[idx["jurusan"]]).strip()
        except IndexError:
            dilewati.append(f"baris {n}: kolom kurang")
            continue
        if not (idp and nama and jurusan):
            dilewati.append(f"baris {n}: ada nilai kosong")
            continue
        if not idp.isdigit():
            dilewati.append(f"baris {n}: id '{idp}' bukan angka")
            continue

        existing = await get_pendaftar(db, idp)
        if existing:
            existing.nama = nama
            existing.jurusan_dituju = jurusan
            diperbarui += 1
        else:
            db.add(Pendaftar(id_pendaftaran=idp, nama=nama, jurusan_dituju=jurusan))
            dibuat += 1
    await db.commit()
    logger.info("Impor master pendaftar: %d baru, %d diperbarui, %d dilewati",
                dibuat, diperbarui, len(dilewati))
    return {"dibuat": dibuat, "diperbarui": diperbarui, "dilewati": dilewati}
