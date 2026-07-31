import csv
import io
import json
import logging

from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.future import select

from models.database import Applicant

logger = logging.getLogger("compliance.export")


def _format_json_list(raw: str | None) -> str:
    """Ubah kolom JSON (fraud_flags/qr_data) menjadi teks yang mudah dibaca."""
    if not raw:
        return ""
    try:
        data = json.loads(raw)
    except (json.JSONDecodeError, TypeError):
        return raw
    if isinstance(data, list):
        return "; ".join(str(f) for f in data)
    return str(data)



def _format_skor_prestasi(raw: str | None) -> str:
    """Ubah JSON skor_prestasi menjadi teks ringkas: total + rincian komponen."""
    if not raw:
        return ""
    try:
        d = json.loads(raw)
    except (json.JSONDecodeError, TypeError):
        return str(raw)
    if not isinstance(d, dict):
        return str(raw)
    total = d.get("total")
    parsial = d.get("total_parsial")
    nilai = total if total is not None else parsial
    label = f"{nilai}" if nilai is not None else "-"
    if total is None and parsial is not None:
        label = f"{parsial} (parsial)"
    komponen = []
    if d.get("bidang"):
        komponen.append(f"{d['bidang']} {d.get('poin_bidang','')}")
    if d.get("tingkat"):
        komponen.append(f"{d['tingkat']} {d.get('poin_tingkat','')}")
    if d.get("partisipasi"):
        komponen.append(f"{d['partisipasi']} {d.get('poin_partisipasi','')}")
    rincian = "; ".join(k.strip() for k in komponen)
    return f"{label} [{rincian}]" if rincian else label


def _bool_label(value) -> str:
    """Ubah nilai boolean/None menjadi label ramah-laporan."""
    if value is True:
        return "Ya"
    if value is False:
        return "Tidak"
    return "-"


# Definisi kolom terpusat -> dipakai CSV maupun XLSX agar selalu konsisten.
HEADERS = [
    "ID", "ID Pendaftaran", "Nama Pendaftar (Form)", "Nama Peserta (Sertifikat)", "Jurusan Tujuan",
    "Nama Lomba", "Singkatan", "Penyelenggara", "Kategori", "Tingkat", "Peringkat",
    "Tanggal Kegiatan", "No. Sertifikat", "URL Verifikasi", "Penandatangan",
    "Ada Cap", "Deskripsi Cap", "Ada TTD",
    "Skor Kepatuhan", "Skor Prestasi", "Status AI", "Status Final",
    "Status Kurasi", "Skor Nama (Kurasi)", "Skor Penyelenggara (Kurasi)", "Ajang Terkurasi Terdekat",
    "Fraud Flags", "QR Data", "Reasoning", "Nama File", "Dibuat",
]


def _row(a: Applicant) -> list:
    """Petakan satu record Applicant ke satu baris laporan (urutan = HEADERS)."""
    return [
        a.id,
        a.id_pendaftaran or "",
        a.applicant_name or "",
        a.nama_peserta or "",
        a.target_major or "",
        a.nama_lomba or "",
        a.singkatan_lomba or "",
        a.nama_penyelenggara or "",
        a.kategori or "",
        a.tingkat or "",
        a.peringkat or "",
        a.tanggal_kegiatan or "",
        a.nomor_sertifikat or "",
        a.url_verifikasi or "",
        a.penandatangan or "",
        _bool_label(a.ada_cap),
        a.deskripsi_cap or "",
        _bool_label(a.ada_ttd),
        a.skor_kepatuhan if a.skor_kepatuhan is not None else "",
        _format_skor_prestasi(a.skor_prestasi),
        a.ai_status or "",
        a.final_status or "",
        a.kurasi_status or "",
        a.kurasi_skor_nama if a.kurasi_skor_nama is not None else "",
        a.kurasi_skor_penyelenggara if a.kurasi_skor_penyelenggara is not None else "",
        a.kurasi_ajang_terdekat or "",
        _format_json_list(a.fraud_flags),
        _format_json_list(a.qr_data),
        a.reasoning or "",
        a.filename or "",
        a.created_at.isoformat() if a.created_at else "",
    ]


async def _fetch_applicants(db: AsyncSession, verifikator_id: int | None = None):
    query = select(Applicant).order_by(Applicant.id.desc())
    if verifikator_id is not None:
        query = query.where(Applicant.verifikator_id == verifikator_id)
    result = await db.execute(query)
    return result.scalars().all()


async def generate_csv_export(
    db: AsyncSession, verifikator_id: int | None = None
) -> io.StringIO:
    """Ekspor pendaftar ke buffer CSV (StringIO)."""
    applicants = await _fetch_applicants(db, verifikator_id)
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(HEADERS)
    for a in applicants:
        writer.writerow(_row(a))
    output.seek(0)
    return output


def _xlsx_safe(value):
    """openpyxl hanya menerima tipe sederhana; sisanya diubah menjadi string."""
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


async def generate_xlsx_export(
    db: AsyncSession, verifikator_id: int | None = None
) -> io.BytesIO:
    """Ekspor pendaftar ke buffer XLSX (BytesIO) memakai openpyxl.

    openpyxl diimpor secara lazy agar aplikasi tetap jalan meski paket belum
    terpasang; error hanya muncul saat fitur XLSX benar-benar dipakai.
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError as exc:
        raise RuntimeError(
            "Paket 'openpyxl' belum terpasang. Jalankan: pip install openpyxl"
        ) from exc

    applicants = await _fetch_applicants(db, verifikator_id)

    wb = Workbook()
    ws = wb.active
    ws.title = "Pendaftar"

    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for a in applicants:
        ws.append([_xlsx_safe(v) for v in _row(a)])

    # Lebar kolom sederhana berdasarkan panjang header (dibatasi 12..45).
    for idx, header in enumerate(HEADERS, start=1):
        letter = ws.cell(row=1, column=idx).column_letter
        ws.column_dimensions[letter].width = min(max(len(header) + 2, 12), 45)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer